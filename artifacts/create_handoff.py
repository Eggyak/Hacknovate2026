from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE

out = r'D:\hacknovate2026\artifacts\MHA Grievance Workspace Handoff.docx'
doc = Document()
sec = doc.sections[0]
sec.top_margin=Inches(.68); sec.bottom_margin=Inches(.62); sec.left_margin=Inches(.72); sec.right_margin=Inches(.72)
styles=doc.styles
styles['Normal'].font.name='Aptos'; styles['Normal'].font.size=Pt(9); styles['Normal'].font.color.rgb=RGBColor(49,65,81)
styles['Normal'].paragraph_format.space_after=Pt(5)
for s,size,color in [('Title',28,'173A59'),('Heading 1',18,'173A59'),('Heading 2',12,'27668C'),('Heading 3',10,'344B61')]:
    st=styles[s]; st.font.name='Aptos Display'; st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(color)
    st.paragraph_format.space_before=Pt(12 if s!='Title' else 0); st.paragraph_format.space_after=Pt(5)

def shade(cell, fill):
    pr=cell._tc.get_or_add_tcPr(); el=OxmlElement('w:shd'); el.set(qn('w:fill'),fill); pr.append(el)
def margins(cell, top=100, start=120, bottom=100, end=120):
    tc=cell._tc; pr=tc.get_or_add_tcPr(); m=OxmlElement('w:tcMar')
    for k,v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node=OxmlElement('w:'+k); node.set(qn('w:w'),str(v)); node.set(qn('w:type'),'dxa'); m.append(node)
    pr.append(m)
def bullet(text): doc.add_paragraph(text,style='List Bullet')
def numbered(text): doc.add_paragraph(text,style='List Number')
def code(text):
    p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.13); p.paragraph_format.space_after=Pt(2)
    r=p.add_run(text); r.font.name='Consolas'; r.font.size=Pt(8); r.font.color.rgb=RGBColor(44,81,108)

# header/footer
header=sec.header.paragraphs[0]; header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=header.add_run('HACKNOVATE 2026  /  PROJECT HANDOFF'); r.font.name='Aptos'; r.font.size=Pt(7); r.font.bold=True; r.font.color.rgb=RGBColor(114,132,149)
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=footer.add_run('MHA Grievance Department Workspace  •  Editable implementation reference'); r.font.size=Pt(7); r.font.color.rgb=RGBColor(128,143,157)

p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(28)
r=p.add_run('PROJECT HANDOFF  /  30 OCTOBER 2026'); r.font.size=Pt(8); r.font.bold=True; r.font.color.rgb=RGBColor(168,129,67)
doc.add_paragraph('MHA Grievance\nDepartment Workspace',style='Title')
p=doc.add_paragraph('Editable reference for the multilingual voice-first grievance dashboard prototype')
p.paragraph_format.space_after=Pt(14); p.runs[0].font.size=Pt(12); p.runs[0].font.color.rgb=RGBColor(99,117,134)

# callout
call=doc.add_table(rows=1,cols=1); call.alignment=WD_TABLE_ALIGNMENT.CENTER; c=call.cell(0,0); shade(c,'EEF4F7'); margins(c,170,200,170,200)
p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0)
r=p.add_run('Current result  '); r.bold=True; r.font.color.rgb=RGBColor(31,91,126)
p.add_run('A runnable React + TypeScript dashboard prototype with email-based demo sign-in with role-scoped access, ward/division filtering, editable role and user management, triage filters, ticket details, and browser-persisted status updates. The interface uses mock data and does not connect to MHA systems or perform real translation/notifications.')

doc.add_heading('1  What is in the workspace',level=1)
rows=[('src/App.tsx','Shared email sign-in, account-specific dashboard, scope-aware inbox, details drawer, officer action desk, role creation and user onboarding modals.'),('src/services/auth.ts','AuthService interface and demo email/password adapter to replace with the backend function.'),('src/styles.css','Responsive visual system for desktop, tablet, and mobile layouts.'),('src/types/grievance.ts','Strict domain types, including role permissions, scope, ticket base fields, and category-specific discriminated payloads.'),('src/data/mockData.ts','Three wards, four demo roles/users, and eight realistic sample grievances.'),('START_MHA_DASHBOARD.bat','Windows launcher: checks Node, installs dependencies on first run, opens the local app and starts Vite.'),('package.json / vite.config.ts','Vite, React, TypeScript and Tailwind v4 project setup.')]
t=doc.add_table(rows=1,cols=2); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Light Shading Accent 1'
t.cell(0,0).text='File'; t.cell(0,1).text='Purpose'
for c in t.rows[0].cells: shade(c,'173A59')
for c in t.rows[0].cells:
    for run in c.paragraphs[0].runs: run.font.bold=True; run.font.color.rgb=RGBColor(255,255,255); run.font.size=Pt(8)
for a,b in rows:
    cells=t.add_row().cells; cells[0].text=a; cells[1].text=b
    for c in cells: margins(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for run in cells[0].paragraphs[0].runs: run.font.name='Consolas'; run.font.size=Pt(8); run.font.color.rgb=RGBColor(43,92,125)
    for run in cells[1].paragraphs[0].runs: run.font.size=Pt(8)

doc.add_page_break()
doc.add_heading('2  Architecture at a glance',level=1)
p=doc.add_paragraph('The prototype is a single-page Vite application served at one shared URL. Sign-in resolves the department email to its assigned user role; that account scope determines which tickets and actions are available. Demo session and workspace changes are stored in this browser.')
# Editable diagram with native Word table cells
flow=doc.add_table(rows=3,cols=3); flow.alignment=WD_TABLE_ALIGNMENT.CENTER; flow.autofit=False
flowPr=flow._tbl.tblPr; flowBorders=OxmlElement('w:tblBorders')
for edge in ('top','left','bottom','right','insideH','insideV'):
    el=OxmlElement('w:'+edge); el.set(qn('w:val'),'nil'); flowBorders.append(el)
flowPr.append(flowBorders)
for row in flow.rows:
    rowPr=row._tr.get_or_add_trPr(); keep=OxmlElement('w:cantSplit'); rowPr.append(keep)
widths=[Inches(2.1),Inches(.55),Inches(3.55)]
for row in flow.rows:
    for i,c in enumerate(row.cells): c.width=widths[i]; margins(c,85,135,85,135)
items=[('ACCOUNT SIGN IN','Department email resolves the assigned user and role','ACCESS RULES','GLOBAL: all tickets\nDIVISION: matching division\nWARD: assigned ward IDs'),('DEMO STATE','Roles · Users · Wards · Tickets','INBOX','Search · status · division filters\nSLA and status indicators'),('LOCAL STORAGE','Role and user edits\nTicket status and ATR','CASE WORKSPACE','Transcript · incident data · location\nOfficer report · simulated response')]
for i,(left,leftsub,right,rightsub) in enumerate(items):
    a,b,c=flow.rows[i].cells
    shade(a,'EAF2F7'); shade(b,'F8F9FA'); shade(c,'F5F7F9')
    a.text=''; p=a.paragraphs[0]; rr=p.add_run(left); rr.bold=True; rr.font.size=Pt(9); rr.font.color.rgb=RGBColor(37,91,124); p.add_run('\n'+leftsub).font.size=Pt(8)
    b.text=''; p=b.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; rr=p.add_run('→'); rr.bold=True; rr.font.size=Pt(15); rr.font.color.rgb=RGBColor(170,132,70)
    c.text=''; p=c.paragraphs[0]; rr=p.add_run(right); rr.bold=True; rr.font.size=Pt(9); rr.font.color.rgb=RGBColor(37,91,124); p.add_run('\n'+rightsub).font.size=Pt(8)

doc.add_heading('3  Demo access and behavior',level=1)
for item in [
'Super Admin (Nodal Officer): GLOBAL scope, all five permissions, creates roles and onboards officers.',
'Cybercrime Desk Officer: DIVISION scope restricted to Cyber & Information Security.',
'Ward 101 Inspector: WARD scope restricted to WARD-101.',
'Multi-Ward Supervisor: WARD scope restricted to WARD-101 and WARD-102.',
'Use Sign out, then sign in with another department email to preview its dashboard at the same URL. The queue, summary cards, SLA overview, and scope label follow the signed-in account.',
'Open a queue row to view transcript, category payload, location, timeline, and the officer action desk. Status and action report persist locally.',
'As Super Admin, create roles with GLOBAL, DIVISION, or WARD scope and choose permissions. Onboard users and assign ward IDs for ward-scoped roles.',
'Role management entry points display a helpful message when the active user lacks users:manage permission.'
]: bullet(item)

doc.add_heading('4  Start the dashboard on Windows',level=1)
numbered('Install Node.js 20 or newer if it is not already installed.')
numbered('Double-click START_MHA_DASHBOARD.bat in the project folder.')
numbered('On first run, the launcher runs npm install. Later runs reuse node_modules.')
numbered('The launcher opens http://localhost:5173 and starts the Vite development server in the command window. Keep that window open while using the dashboard; press Ctrl+C to stop it.')
code('npm run build   # production compile and bundle')

doc.add_heading('5  Schema and mock data',level=1)
p=doc.add_paragraph('The Ticket union ties each category to a typed details object and division ID. Common fields include grievance ID, complainant profile and OTP state, mobile, full address, pincode, GPS coordinates, ward, status, SLA deadline, language, truth declaration, audio URL, regional text, English text, and evidence attachments.')
for item in ['Cybercrime: transaction/UTR, lost amount, incident time, platform, suspect details.','Immigration: passport, application reference, nationality, visa/OCI type, issue and expiry dates.','Disaster relief: masked Aadhaar, bank account and IFSC, disaster, damage category.','UT administration: UT, civic department, ward/zone and property ID.','Additional divisions from the supplied requirements document are represented by typed payloads for freedom fighter rehabilitation, CAPF misconduct, NGO/FCRA, internal security, counter terrorism, LWE/North East, border management, judicial, police, Centre-State, and internal administration.']: bullet(item)
p=doc.add_paragraph('Mock tickets use Delhi wards 101, 102, and 103, Hindi/Tamil/Urdu/English and other language examples, and four requested divisions. Personal data values are fictional and partially masked where sensitive fields are represented.')

doc.add_heading('6  Prototype boundaries and next engineering work',level=1)
for item in [
'The email/password flow is demo authentication implemented in the browser. It is not secure, does not provide server-side authorization, and must not be used for real accounts. Replace src/services/auth.ts with the team authentication API and enforce identity, permissions, and scope on the server. There is no audit store, OTP provider, real SLA scheduler, or production database.',
'Audio uses a local illustrative demo clip. Replace audioFileUrl with each real intake recording URL when the citizen frontend/backend is integrated.',
'Translation and citizen notifications are simulated. Connect approved translation, speech, and notification services before representing them as live capabilities.',
'Browser localStorage is for a single-device demo only; it is not secure, multi-user, or suitable for sensitive production grievance data.',
'Next: agree API contracts with the chatbot team; build authenticated backend endpoints; enforce every scope and permission on the server; add audit history and evidence storage; implement SLA calculations from deadlines; validate schemas at runtime; add tests and accessibility review.',
'Before a public demo, replace fictional data and local URLs as needed, and clear the demo browser storage using DevTools Application → Local Storage → localhost:5173.'
]: bullet(item)

doc.add_heading('7  Files to continue with',level=1)
for item in ['Data shape or categories: edit src/types/grievance.ts, then update src/data/mockData.ts.','Workspace interactions and UI: edit src/App.tsx. Replace the demo auth adapter in src/services/auth.ts with the team login function.','Visual styling and responsive breakpoints: edit src/styles.css.','Run the prototype: START_MHA_DASHBOARD.bat.','Confirm the code compiles: npm run build.']: bullet(item)

# properties
doc.core_properties.title='MHA Grievance Department Workspace Handoff'
doc.core_properties.subject='Editable implementation reference for the Hacknovate 2026 MHA dashboard prototype'
doc.core_properties.author='Hacknovate 2026 project team'
doc.save(out)
print(out)




